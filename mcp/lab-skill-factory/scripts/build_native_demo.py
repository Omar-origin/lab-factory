#!/usr/bin/env python3
"""Rebuild the public self-authored example using the shipped native engine."""

import hashlib
import json
import sys
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "skills/lab-skill-factory/scripts"))
import v2_engine as engine


def build():
    workspace = ROOT / "reports/implementation-2026-09-13/demo"
    workspace.mkdir(parents=True, exist_ok=True)
    image = ROOT / "mcp/lab-skill-factory/evals/native-website-screenshot.png"
    target_image = workspace / "public-docs-screenshot.png"
    target_image.write_bytes(image.read_bytes())
    source = workspace / "native-template.docx"
    doc = Document()
    doc.add_heading("Lab Factory 原生图表演示", 0)
    doc.add_paragraph(
        "自建功能样稿：下表为演示数据，图片为官网页面截图，不是学生实验结果。"
    )
    doc.add_paragraph("<演示内容待填>")
    doc.save(source)
    inv = engine.inventory_docx(source)
    selected = next(n for n in inv["nodes"] if n.get("text") == "<演示内容待填>")
    profile = engine.create_template_profile(
        inv,
        [
            {
                "id": "body",
                "node_id": selected["node_id"],
                "operation": "replace_placeholder",
            }
        ],
        "功能演示",
    )
    package = {
        "version": "3.0",
        "chapter": "1",
        "assets": {
            "website": {
                "relative_path": target_image.name,
                "sha256": hashlib.sha256(target_image.read_bytes()).hexdigest(),
            }
        },
        "items": [
            {
                "field_id": "body",
                "blocks": [
                    {
                        "type": "paragraph",
                        "id": "intro",
                        "segments": [
                            {"text": "演示对象及其编辑方式列于"},
                            {"ref": "features"},
                            {"text": "。这些内容用于检查文件结构。"},
                        ],
                    },
                    {
                        "type": "table",
                        "id": "features",
                        "caption_id": "features",
                        "caption": "演示对象说明",
                        "columns": ["对象", "保存方式", "检查方式"],
                        "rows": [
                            ["截图", "DOCX 内嵌图片", "点击图片并检查比例"],
                            ["三线表", "Word 原生表格", "修改单元格并保存"],
                            ["图表编号", "正文引用与题注", "检查编号对应关系"],
                        ],
                    },
                    {
                        "type": "paragraph",
                        "id": "image_intro",
                        "segments": [
                            {"text": "审查时截取的官网文档页面见"},
                            {"ref": "website"},
                            {
                                "text": "。该图片仅作为插图演示，不用于证明任何实验已完成。"
                            },
                        ],
                    },
                    {
                        "type": "figure",
                        "id": "website",
                        "asset_id": "website",
                        "caption_id": "website",
                        "caption": "官网文档页截图（演示材料）",
                    },
                ],
            }
        ],
    }
    engine.write_json(workspace / "content-package.json", package)
    result = engine.apply_v2(
        profile,
        source,
        package,
        workspace / "native-report.docx",
        overwrite=True,
        workspace=workspace,
    )
    engine.write_json(
        workspace / "validation.json",
        {
            "writeback": result,
            "structure": engine.document_structure_audit(
                workspace / "native-report.docx"
            ),
        },
    )
    public = ROOT / "mcp/lab-skill-factory/cloudflare/public/examples"
    public.mkdir(exist_ok=True)
    for name in ("native-template.docx", "native-report.docx"):
        (public / name).write_bytes((workspace / name).read_bytes())
    print(
        json.dumps(
            {"ok": True, "output": str(workspace / "native-report.docx")},
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    build()
