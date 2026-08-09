#!/usr/bin/env python3
"""Run 20 beta smoke tests for fill-map writeback."""

from __future__ import annotations

import hashlib
import json
import subprocess
import sys
import tempfile
from pathlib import Path


HERE = Path(__file__).resolve()
BUNDLED_ROOT = HERE.parents[1]
if (BUNDLED_ROOT / "evals" / "beta-20-samples.json").exists() and (
    BUNDLED_ROOT / "skills" / "lab-skill-factory"
).exists():
    ROOT = BUNDLED_ROOT
    MCP_DIR = BUNDLED_ROOT
    SKILL_DIR = BUNDLED_ROOT / "skills" / "lab-skill-factory"
else:
    ROOT = HERE.parents[3]
    MCP_DIR = ROOT / "mcp" / "lab-skill-factory"
    SKILL_DIR = ROOT / "skills" / "lab-skill-factory"
SAMPLES = MCP_DIR / "evals" / "beta-20-samples.json"
APPLY = SKILL_DIR / "scripts" / "apply_fill_map.py"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def make_fill(path: Path, heading: str, content: str) -> None:
    path.write_text(f"# 填补内容预览\n\n## {heading}\n\n{content}\n", encoding="utf-8")


def make_md_target(path: Path, sample: dict) -> None:
    if sample["operation"] == "replace_placeholder":
        body = f"# 测试报告\n\n## 正文内容\n\n{sample['anchor']}\n"
    elif sample["operation"] == "fill_table_cell":
        body = f"# 测试报告\n\n| 项目 | 内容 |\n| --- | --- |\n| A | {sample['anchor']} |\n"
    else:
        body = (
            "# 测试报告\n\n"
            "## 实验目的\n\n"
            "## 实验环境\n\n"
            "## 实验截图\n\n"
            "## 实验小结\n\n"
            "## 附录\n\n"
            "## 源码说明\n\n"
            "## 参考案例\n\n"
        )
    path.write_text(body, encoding="utf-8")


def make_docx_target(path: Path, sample: dict) -> None:
    from docx import Document

    document = Document()
    document.add_heading("测试报告", level=1)
    for anchor in ["实验目的", "实验环境", "实验截图", "实验小结", "附录", "参考案例"]:
        document.add_paragraph(anchor)
    document.add_paragraph("<待填写正文>")
    document.add_paragraph("<工具环境>")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "内容"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "<表格待填>"
    document.save(str(path))


def docx_text(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def run_sample(sample: dict, workspace: Path) -> dict:
    sample_dir = workspace / sample["id"]
    sample_dir.mkdir(parents=True)
    fill = sample_dir / "fill.md"
    target = sample_dir / ("target.docx" if sample["format"] == "docx" else "target.md")
    output = sample_dir / ("draft.docx" if sample["format"] == "docx" else "draft.md")
    fill_map = sample_dir / "fill-map.json"

    make_fill(fill, sample["heading"], sample["content"])
    if sample["format"] == "docx":
        make_docx_target(target, sample)
    else:
        make_md_target(target, sample)
    source_hash = sha256(target)
    fill_map.write_text(
        json.dumps(
            {
                "version": "0.1.0",
                "source_fill": "fill.md",
                "target_document": target.name,
                "copy_mode": "byte_for_byte_first",
                "format_strategy": "inherit_target_anchor",
                "items": [
                    {
                        "id": sample["id"],
                        "source": f"fill.md#{sample['heading']}",
                        "target_anchor": sample["anchor"],
                        "operation": sample["operation"],
                        "preserve_original": True,
                        "format_strategy": "inherit_target_anchor",
                    }
                ],
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    if getattr(sys, "frozen", False):
        # The built-in test command operates only on synthetic temporary files;
        # execute the bundled helper directly so testing does not require a paid license.
        command = [sys.executable, "--run-skill-script", str(APPLY), str(fill_map), "--output", str(output)]
    else:
        command = [sys.executable, str(APPLY), str(fill_map), "--output", str(output)]
    proc = subprocess.run(
        command,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    contains = False
    if output.exists():
        text = docx_text(output) if sample["format"] == "docx" else output.read_text(encoding="utf-8")
        contains = sample["content"] in text
    return {
        "id": sample["id"],
        "ok": proc.returncode == 0 and output.exists() and contains and sha256(target) == source_hash,
        "returncode": proc.returncode,
        "source_unchanged": sha256(target) == source_hash,
        "output_exists": output.exists(),
        "content_found": contains,
        "stdout": proc.stdout.strip(),
        "stderr": proc.stderr.strip(),
    }


def main() -> int:
    data = json.loads(SAMPLES.read_text(encoding="utf-8"))
    samples = data["samples"]
    with tempfile.TemporaryDirectory(prefix="lab-factory-beta-smoke-") as tmp:
        workspace = Path(tmp)
        results = [run_sample(sample, workspace) for sample in samples]
    failed = [result for result in results if not result["ok"]]
    print(json.dumps({"ok": not failed, "count": len(results), "failed": failed, "results": results}, ensure_ascii=False, indent=2))
    return 0 if not failed else 1


if __name__ == "__main__":
    raise SystemExit(main())
