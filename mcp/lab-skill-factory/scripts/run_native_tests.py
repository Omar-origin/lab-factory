#!/usr/bin/env python3
"""Native writeback integration and audit trust regressions; uses no provider credentials."""

import copy
import hashlib
import importlib.util
import json
import os
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "mcp/lab-skill-factory"))
sys.path.insert(0, str(ROOT / "skills/lab-skill-factory/scripts"))
os.environ["LAB_FACTORY_DEV_ALLOW"] = "1"
import v2_engine as engine
import artifact_store
import autopilot
import server
from unittest.mock import patch


class NativeTests(unittest.TestCase):
    def setUp(self):
        identity = patch.object(
            server.commercial_client,
            "install_id",
            return_value="native-test-installation",
        )
        identity.start()
        self.addCleanup(identity.stop)
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.source = self.root / "template.docx"
        doc = Document()
        doc.add_paragraph("<正文待填>")
        doc.save(self.source)
        inv = engine.inventory_docx(self.source)
        location = next(x for x in inv["nodes"] if x.get("text") == "<正文待填>")
        self.profile = engine.create_template_profile(
            inv,
            [
                {
                    "id": "body",
                    "node_id": location["node_id"],
                    "operation": "replace_placeholder",
                }
            ],
            "演示",
        )
        self.package = {
            "version": "3.0",
            "items": [
                {
                    "field_id": "body",
                    "blocks": [
                        {
                            "type": "paragraph",
                            "id": "explain",
                            "segments": [
                                {"text": "演示字段及其含义列于"},
                                {"ref": "fields"},
                                {"text": "。"},
                            ],
                        },
                        {
                            "type": "table",
                            "id": "table",
                            "caption_id": "fields",
                            "caption": "演示字段说明",
                            "columns": ["字段", "含义"],
                            "rows": [["name", "示例名称"], ["count", "示例数量"]],
                        },
                    ],
                }
            ],
        }
        self.output = self.root / "draft.docx"

    def tearDown(self):
        self.tmp.cleanup()

    def apply(self, package=None):
        return engine.apply_v2(
            self.profile,
            self.source,
            package or self.package,
            self.output,
            workspace=self.root,
        )

    def test_native_table_and_preservation(self):
        before = engine.zip_entry_hashes(self.source)
        self.apply()
        doc = Document(self.output)
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual(doc.tables[0].cell(1, 0).text, "name")
        after = engine.zip_entry_hashes(self.output)
        self.assertTrue(
            all(
                after[n] == sha for n, sha in before.items() if n != "word/document.xml"
            )
        )
        audit = engine.document_structure_audit(self.output)
        self.assertEqual(audit["status"], "pass", audit)
        self.assertEqual(audit["pending_evidence_count"], 0)
        with zipfile.ZipFile(self.output) as z:
            root = etree.fromstring(z.read("word/document.xml"))
            ns = engine.NS
            self.assertEqual(
                root.xpath(
                    'count(//w:tblPr/w:tblBorders/w:insideV[@w:val="nil"])',
                    namespaces=ns,
                ),
                1,
            )
            self.assertEqual(
                root.xpath("count(//w:trPr/w:tblHeader)", namespaces=ns), 1
            )

    def test_image_hash_and_package(self):
        # Real screenshot captured for the earlier public-page audit; labeled as a demonstration.
        image = ROOT / "mcp/lab-skill-factory/evals/native-website-screenshot.png"
        asset = server.tool_v3_import_asset(
            {
                "workspace": str(self.root),
                "source_path": str(image),
                "source_description": "官网演示截图",
            }
        )
        package = copy.deepcopy(self.package)
        package["assets"] = {asset["asset_id"]: asset["asset"]}
        package["items"][0]["blocks"] += [
            {
                "type": "paragraph",
                "id": "figure_ref",
                "segments": [
                    {"text": "官网演示截图见"},
                    {"ref": "screen"},
                    {"text": "。"},
                ],
            },
            {
                "type": "figure",
                "id": "screen",
                "asset_id": asset["asset_id"],
                "caption_id": "screen",
                "caption": "官网演示截图，非实验结果",
            },
        ]
        result = self.apply(package)
        self.assertEqual(result["native_objects"]["image_count"], 1)
        self.assertEqual(engine.document_structure_audit(self.output)["status"], "pass")
        package["assets"][asset["asset_id"]]["sha256"] = "0" * 64
        self.output.unlink()
        with self.assertRaises(engine.V2Error):
            self.apply(package)
        self.assertFalse(self.output.exists())

    def test_missing_cross_reference_rejected_transactionally(self):
        self.package["items"][0]["blocks"].pop(0)
        with self.assertRaises(engine.V2Error):
            self.apply()
        self.assertFalse(self.output.exists())

    def test_audit_and_mutation_binding(self):
        self.apply()
        ledger = self.root / "facts.json"
        ledger.write_text(
            json.dumps({"facts": [{"source": "demonstration", "status": "provided"}]})
        )
        result = server.tool_v3_verify_draft(
            {
                "workspace": str(self.root),
                "document_path": str(self.output),
                "fact_ledger_path": str(ledger),
            }
        )
        self.assertTrue(result["ok"], result)
        self.assertEqual(
            autopilot.validated_quality_gates({"draft": result["draft"]}, self.root)[
                "status"
            ],
            "pass",
        )
        with self.output.open("ab") as f:
            f.write(b"changed")
        self.assertEqual(
            autopilot.validated_quality_gates({"draft": result["draft"]}, self.root)[
                "status"
            ],
            "retryable_failure",
        )

    def test_fabricated_checks_rejected(self):
        gates = autopilot.validated_quality_gates(
            {
                "draft": {"object_id": "invented", "sha256": "a" * 64},
                "quality_gates": {"status": "pass"},
            },
            self.root,
        )
        self.assertEqual(gates["status"], "retryable_failure")

    def test_structure_alone_is_notice(self):
        docs = []
        for index, char in enumerate(["甲", "乙"]):
            path = self.root / f"{index}.docx"
            doc = Document()
            for number in range(8):
                doc.add_paragraph(char * 20 + str(number))
            doc.save(path)
            docs.append(path)
        result = engine.cohort_similarity_check(docs[0], [docs[1]], [])
        self.assertEqual(result["gate"], "pass", result)
        self.assertTrue(result["findings"][0]["structure_notice_only"])

    def test_invalid_prepare_does_not_consume(self):
        with patch.object(server, "consume_report_use") as consume:
            with self.assertRaises(Exception):
                server.tool_v2_prepare_autopilot(
                    {"idempotency_key": "00000000-0000-4000-8000-000000000001"}
                )
            consume.assert_not_called()

    def test_pending_evidence_cannot_be_certified(self):
        doc = Document()
        doc.add_paragraph("如图 1-1 所示。")
        doc.add_paragraph("【图 1-1：运行截图；待补：真实截图】")
        doc.save(self.output)
        engine.ensure_docx_disclaimer(self.output)
        ledger = self.root / "facts.json"
        ledger.write_text(
            json.dumps({"facts": [{"source": "user", "status": "unknown"}]})
        )
        result = server.tool_v3_verify_draft(
            {
                "workspace": str(self.root),
                "document_path": str(self.output),
                "fact_ledger_path": str(ledger),
            }
        )
        self.assertFalse(result["ok"])

    def test_usage_timeout_replays_same_report(self):
        import uuid

        requirements = self.root / "requirements.json"
        requirements.write_text(
            json.dumps(
                {
                    "version": "2.0",
                    "experiment": {},
                    "formatting": {},
                    "missing": [],
                    "conflicts": [],
                }
            )
        )
        profile = self.root / "profile.json"
        profile.write_text(json.dumps(self.profile))
        args = {
            "workspace": str(self.root),
            "subject": "测试",
            "idempotency_key": str(uuid.uuid4()),
            "requirements_path": str(requirements),
            "template_profile_path": str(profile),
        }
        keys = []

        def consume(report_id):
            keys.append(report_id)
            if len(keys) == 1:
                raise RuntimeError("simulated response timeout after remote commit")
            return {"ok": True}

        with patch.object(server, "consume_report_use", side_effect=consume):
            with self.assertRaises(RuntimeError):
                server.tool_v2_prepare_autopilot(args)
            replay = server.tool_v2_prepare_autopilot(args)
        self.assertEqual(len(set(keys)), 1)
        self.assertEqual(keys[0], replay["session"]["report_id"])
        receipts = list((self.root / "lab-factory/usage").glob("*.json"))
        self.assertEqual(json.loads(receipts[0].read_text())["status"], "acknowledged")

    def test_balanced_native_end_to_end_and_changed_final(self):
        import uuid

        key = lambda: str(uuid.uuid4())
        requirements = self.root / "requirements.json"
        requirements.write_text(
            json.dumps(
                {
                    "version": "2.0",
                    "experiment": {},
                    "formatting": {},
                    "missing": [],
                    "conflicts": [],
                    "required_artifacts": [
                        {"id": "fields", "kind": "table", "source": "test task"}
                    ],
                }
            )
        )
        profile = self.root / "profile.json"
        profile.write_text(json.dumps(self.profile))
        self.package["requirement_coverage"] = {"fields": "table"}
        package_path = self.root / "content.json"
        package_path.write_text(json.dumps(self.package))
        prepared = server.tool_v2_prepare_autopilot(
            {
                "workspace": str(self.root),
                "subject": "测试",
                "idempotency_key": key(),
                "requirements_path": str(requirements),
                "template_profile_path": str(profile),
            }
        )
        checkpoint = prepared["interaction_plan"]["checkpoint"]
        confirmed = autopilot.confirm_checkpoint(
            self.root,
            prepared["session"]["state_version"],
            key(),
            "preflight",
            checkpoint["summary_sha256"],
        )
        running = autopilot.advance(
            self.root,
            confirmed["session"]["state_version"],
            key(),
            confirmed["confirmation_token"],
            {},
        )
        result = server.tool_v3_apply_draft(
            {
                "workspace": str(self.root),
                "profile_path": str(profile),
                "docx_path": str(self.source),
                "content_package_path": str(package_path),
                "output_path": str(self.output),
            }
        )
        self.assertEqual(result["requirement_coverage"]["checked_count"], 1)
        ledger = self.root / "facts.json"
        ledger.write_text(
            json.dumps({"facts": [{"source": "test task", "status": "provided"}]})
        )
        verified = server.tool_v3_verify_draft(
            {
                "workspace": str(self.root),
                "document_path": str(self.output),
                "fact_ledger_path": str(ledger),
            }
        )
        self.assertTrue(verified["ok"], verified)
        preview = autopilot.advance(
            self.root,
            running["session"]["state_version"],
            key(),
            None,
            {"draft": verified["draft"]},
        )
        checkpoint = preview["interaction_plan"]["checkpoint"]
        original = self.output.read_bytes()
        self.output.write_bytes(original + b"changed")
        with self.assertRaises(autopilot.AutopilotError):
            autopilot.confirm_checkpoint(
                self.root,
                preview["session"]["state_version"],
                key(),
                "final_review",
                checkpoint["summary_sha256"],
            )
        self.output.write_bytes(original)
        confirmed = autopilot.confirm_checkpoint(
            self.root,
            preview["session"]["state_version"],
            key(),
            "final_review",
            checkpoint["summary_sha256"],
        )
        self.output.write_bytes(original + b"changed")
        with self.assertRaises(autopilot.AutopilotError):
            autopilot.advance(
                self.root,
                confirmed["session"]["state_version"],
                key(),
                confirmed["confirmation_token"],
                {"action": "accept"},
            )
        self.output.write_bytes(original)
        finalized = autopilot.advance(
            self.root,
            confirmed["session"]["state_version"],
            key(),
            confirmed["confirmation_token"],
            {"action": "accept"},
        )
        self.assertEqual(finalized["session"]["orchestration_state"], "finalized")

    def test_native_table_in_template_cell(self):
        doc = Document()
        outer = doc.add_table(rows=1, cols=1)
        outer.cell(0, 0).text = "<正文待填>"
        doc.save(self.source)
        inv = engine.inventory_docx(self.source)
        location = next(x for x in inv["nodes"] if x.get("text") == "<正文待填>")
        self.profile = engine.create_template_profile(
            inv,
            [{"id": "body", "node_id": location["node_id"], "operation": "fill_cell"}],
            "演示",
        )
        self.apply()
        self.assertEqual(len(Document(self.output).tables[0].cell(0, 0).tables), 1)

    def test_invalid_native_matrix_rejected(self):
        self.package["items"][0]["blocks"][1]["rows"] = [["missing column"]]
        with self.assertRaises(engine.V2Error):
            self.apply()
        self.assertFalse(self.output.exists())

    def test_same_template_and_table_do_not_count_as_authored_prose(self):
        doc = Document()
        doc.add_paragraph("课程固定说明：" + "请依据实验记录填写报告并核对格式。" * 10)
        doc.add_paragraph("<正文待填>")
        doc.save(self.source)
        inv = engine.inventory_docx(self.source)
        location = next(x for x in inv["nodes"] if x.get("text") == "<正文待填>")
        self.profile = engine.create_template_profile(
            inv,
            [
                {
                    "id": "body",
                    "node_id": location["node_id"],
                    "operation": "replace_placeholder",
                }
            ],
            "演示",
        )
        outputs = []
        for index, char in enumerate(["甲", "乙"]):
            package = copy.deepcopy(self.package)
            package["items"][0]["blocks"][0]["segments"] = [
                {"text": char * 100},
                {"ref": "fields"},
            ]
            self.output = self.root / f"scoped-{index}.docx"
            self.apply(package)
            outputs.append(self.output)
        result = engine.cohort_similarity_check(outputs[0], [outputs[1]], [])
        self.assertEqual(result["gate"], "pass", result)


if __name__ == "__main__":
    unittest.main(verbosity=2)
