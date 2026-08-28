#!/usr/bin/env python3
"""Self-contained regression tests for the Lab Factory v2 engine."""

from __future__ import annotations

import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


SCRIPT_ROOT = Path(__file__).resolve().parents[1]
ROOT = SCRIPT_ROOT if (SCRIPT_ROOT / "skills" / "lab-skill-factory").exists() else Path(__file__).resolve().parents[3]
ENGINE_PATH = ROOT / "skills" / "lab-skill-factory" / "scripts" / "v2_engine.py"
SPEC = importlib.util.spec_from_file_location("lab_factory_v2_engine", ENGINE_PATH)
assert SPEC and SPEC.loader
ENGINE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(ENGINE)
GOLDEN_FORMAT = SCRIPT_ROOT / "evals" / "docx-golden-format.json"


def make_template(path: Path) -> None:
    document = Document()
    document.add_heading("实验报告", level=1)
    document.add_paragraph("实验目的")
    document.add_paragraph("请在此处填写实验目的")
    table = document.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "栏目"
    table.cell(0, 1).text = "填写内容"
    table.cell(1, 0).text = "实验结果"
    table.cell(1, 1).text = "<结果待填>"
    table.cell(2, 0).text = "实验小结"
    table.cell(2, 1).text = "<小结待填>"
    document.add_paragraph("实验结果")
    document.add_paragraph("此处仅为提交说明，不应写入")
    if not getattr(sys, "frozen", False):
        section = document.sections[0]
        section.header.paragraphs[0].text = "课程实验报告"
    document.save(path)


def make_formatting_template(path: Path) -> None:
    document = Document()
    prompt = document.add_paragraph()
    prompt_run = prompt.add_run("<正文待填>")
    prompt_run.font.name = "宋体"
    prompt_run.font.size = Pt(12)
    prompt_run.font.color.rgb = RGBColor(255, 0, 0)

    heading = document.add_paragraph()
    heading_run = heading.add_run("实验小结")
    heading_run.bold = True
    heading_run.font.size = Pt(15)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "3")
    num_pr.append(ilvl)
    num_pr.append(num_id)
    heading._p.get_or_add_pPr().insert(0, num_pr)

    body = document.add_paragraph()
    body_run = body.add_run("")
    body_run.font.name = "宋体"
    body_run.font.size = Pt(12)
    body_run.font.color.rgb = RGBColor(0, 0, 0)
    document.add_paragraph("2.2.2 XXX功能")
    document.add_paragraph("3.3 XXX待完善")
    document.save(path)


def make_section_template(path: Path) -> None:
    document = Document()
    document.add_heading("1 系统设计", level=1)
    document.add_heading("1.1 功能设计", level=2)
    document.add_heading("1.1.1 登录功能", level=3)
    document.add_paragraph("登录功能正文。")
    document.save(path)


def make_structured_report(path: Path, theme: str, valid: bool = True) -> None:
    document = Document()
    document.add_heading(f"{theme}实验报告", level=1)
    document.add_paragraph(f"本节先说明{theme}的操作目标和验证边界，随后使用真实证据位置辅助说明。")
    document.add_paragraph(f"第一步完成{theme}环境检查，并记录输入条件、操作顺序和预期结果。")
    if valid:
        document.add_paragraph("如图 2-1 所示，该界面用于核对关键状态。")
        document.add_paragraph(f"【图 2-1：{theme}关键状态界面；待补：真实运行截图】")
        document.add_paragraph("核心字段和职责见表 2-1，真实值由后续运行结果补充。")
        document.add_paragraph(f"【表 2-1：{theme}核心字段与职责；待补：字段、类型、约束和来源】")
    else:
        document.add_paragraph("2.2.1 XXX功能")
        document.add_paragraph("【绘图：请同时补充类图、活动图和状态图】")
    document.add_paragraph(f"第二步围绕{theme}的异常分支进行检查，重点比较预期行为和实际反馈。")
    document.add_paragraph(f"第三步整理{theme}的持久化约束，避免把页面参数直接写入数据层。")
    document.add_paragraph("问题1：输入条件不完整。解决办法：先校验必要字段，再执行后续步骤。")
    document.add_paragraph("心得体会：本次实验让我更清楚地看到验证顺序对定位问题的影响。")
    document.add_paragraph("建议：后续增加边界输入和恢复路径的真实测试记录。")
    document.save(path)


def make_problem_evidence_report(path: Path, include_placeholder: bool) -> None:
    document = Document()
    document.add_heading("异常处理实验报告", level=1)
    document.add_paragraph("本次实验先完成环境配置，再使用边界输入检查异常路径。")
    document.add_paragraph("问题和解决办法")
    document.add_paragraph("问题1：运行时出现错误提示，无法从文字描述判断具体配置状态。")
    if include_placeholder:
        document.add_paragraph("解决过程中的界面状态如图 4-1 所示。")
        document.add_paragraph("【图 4-1：异常配置提示界面；待补：包含真实错误信息的运行截图】")
    document.add_paragraph("解决办法：核对配置项后重新运行，并保留修复前后的验证记录。")
    document.save(path)


def xml_entry_hashes(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        return {name: ENGINE.sha256(archive.read(name)) for name in archive.namelist()}


def add_alternate_text_box(path: Path, text_value: str = "兼容文本框") -> None:
    with zipfile.ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    root = ENGINE.etree.fromstring(entries["word/document.xml"])
    body = root.find(f".//{ENGINE.W}body")
    alternate = ENGINE.etree.Element(f"{{{ENGINE.NS['mc']}}}AlternateContent")
    choice = ENGINE.etree.SubElement(alternate, f"{{{ENGINE.NS['mc']}}}Choice")
    choice.set("Requires", "wps")
    fallback = ENGINE.etree.SubElement(alternate, f"{{{ENGINE.NS['mc']}}}Fallback")
    for branch in (choice, fallback):
        text_box = ENGINE.etree.SubElement(branch, f"{ENGINE.W}txbxContent")
        paragraph = ENGINE.etree.SubElement(text_box, f"{ENGINE.W}p")
        run = ENGINE.etree.SubElement(paragraph, f"{ENGINE.W}r")
        text = ENGINE.etree.SubElement(run, f"{ENGINE.W}t")
        text.text = text_value
    section_properties = body.find(f"{ENGINE.W}sectPr")
    if section_properties is None:
        body.append(alternate)
    else:
        body.insert(body.index(section_properties), alternate)
    entries["word/document.xml"] = ENGINE.etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
    rewritten = path.with_suffix(".rewritten.docx")
    with zipfile.ZipFile(rewritten, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, value in entries.items():
            archive.writestr(name, value)
    os.replace(rewritten, path)


def run() -> dict:
    results = []
    with tempfile.TemporaryDirectory(prefix="lab-factory-v2-tests-") as temporary_dir:
        root = Path(temporary_dir)
        template = root / "template.docx"
        output = root / "draft.docx"
        make_template(template)
        source_hash = ENGINE.file_sha256(template)

        inventory = ENGINE.inventory_docx(template)
        header_expected = not getattr(sys, "frozen", False)
        results.append({"name": "inventory-stories", "ok": (inventory["story_counts"].get("header1") == 1) == header_expected})
        result_node = next(
            node for node in inventory["nodes"]
            if node["container"] == "table_cell" and node["text"] == "<结果待填>"
        )
        summary_node = next(
            node for node in inventory["nodes"]
            if node["container"] == "table_cell" and node["text"] == "<小结待填>"
        )
        results.append({"name": "inventory-table-coordinate", "ok": result_node["coordinates"] == {"table": 0, "row": 1, "cell": 1, "paragraph": 0}})

        alternate_template = root / "alternate-template.docx"
        make_template(alternate_template)
        add_alternate_text_box(alternate_template)
        alternate_inventory = ENGINE.inventory_docx(alternate_template)
        alternate_nodes = [node for node in alternate_inventory["nodes"] if node["text"] == "兼容文本框"]
        results.append({
            "name": "alternate-content-deduplicated",
            "ok": len(alternate_nodes) == 1 and alternate_nodes[0]["unsupported"] == ["text_box"],
        })
        alternate_cue_template = root / "alternate-cue-template.docx"
        make_template(alternate_cue_template)
        add_alternate_text_box(alternate_cue_template, "兼容 XXX 文本框")
        alternate_cues = ENGINE.remaining_template_cues(alternate_cue_template)
        results.append({
            "name": "unsupported-template-cue-manual-only",
            "ok": any(
                cue["manual_only"] and "text_box" in cue["unsupported"]
                for cue in alternate_cues
                if "XXX" in cue["text"]
            ),
        })

        profile = ENGINE.create_template_profile(
            inventory,
            [
                {"id": "result", "node_id": result_node["node_id"], "operation": "replace_placeholder", "semantic_role": "实验结果"},
                {"id": "summary", "node_id": summary_node["node_id"], "operation": "fill_cell", "semantic_role": "实验小结"},
            ],
            "软件工程",
        )
        proposals = ENGINE.propose_placements(profile, inventory)
        results.append({"name": "placements-auto", "ok": proposals["summary"] == {"auto": 2, "confirm": 0, "blocked": 0}})
        related = root / "related.docx"
        related_document = Document(template)
        related_document.add_paragraph("本次实验新增的长段落内容不会改变模板家族判断。")
        related_document.add_paragraph("另一段填写后的实验说明。")
        related_document.save(related)
        related_inventory = ENGINE.inventory_docx(related)
        related_plan = ENGINE.propose_placements(profile, related_inventory)
        results.append({
            "name": "family-compatible-after-content-growth",
            "ok": (
                not related_plan["family_fingerprint_match"]
                and related_plan["family_compatible"]
                and related_plan["family_compatibility"]["score"] >= ENGINE.FAMILY_COMPATIBILITY_THRESHOLD
                and related_plan["summary"]["auto"] == 2
            ),
        })
        confirmed = ENGINE.confirmed_profile(profile, inventory, proposals, [], "用户确认当前模板位置")
        results.append({"name": "profile-drift-history", "ok": len(confirmed["drift_history"]) == 1})

        incompatible = root / "incompatible.docx"
        incompatible_document = Document()
        incompatible_document.add_paragraph("实验结果")
        incompatible_document.add_paragraph("这里没有表格")
        incompatible_document.save(incompatible)
        incompatible_plan = ENGINE.propose_placements(profile, ENGINE.inventory_docx(incompatible))
        results.append({"name": "container-conflict-blocked", "ok": incompatible_plan["summary"]["blocked"] == 2})
        results.append({"name": "unrelated-family-rejected", "ok": incompatible_plan["family_compatible"] is False})

        content = {
            "version": "2.0",
            "items": [
                {"field_id": "result", "content": "程序运行成功，输出与预期一致。"},
                {"field_id": "summary", "content": "本次实验重点验证了表格内安全写入。"},
            ],
        }
        before_entries = xml_entry_hashes(template)
        apply_result = ENGINE.apply_v2(profile, template, content, output)
        after_entries = xml_entry_hashes(output)
        untouched = all(
            before_entries[name] == after_entries[name]
            for name in before_entries
            if name not in apply_result["changed_parts"]
        )
        output_inventory = ENGINE.inventory_docx(output)
        output_text = "\n".join(node["text"] for node in output_inventory["nodes"])
        with zipfile.ZipFile(output) as archive:
            output_root = ENGINE.etree.fromstring(archive.read("word/document.xml"))
        output_body = output_root.xpath("/w:document/w:body", namespaces=ENGINE.NS)[0]
        disclosure = output_body[0]
        results.append({"name": "source-unchanged", "ok": ENGINE.file_sha256(template) == source_hash})
        results.append({"name": "package-preservation", "ok": untouched and apply_result["changed_parts"] == ["word/document.xml"]})
        results.append({"name": "table-writeback", "ok": "程序运行成功" in output_text and "表格内安全写入" in output_text})
        results.append({
            "name": "disclaimer-visible-at-document-start",
            "ok": (
                ENGINE.element_text(disclosure) == ENGINE.DISCLAIMER_TEXT
                and ENGINE.child_value(disclosure, "./w:pPr/w:jc") == "center"
                and ENGINE.child_value(disclosure, ".//w:rPr/w:color") == "C00000"
                and ENGINE.child_value(disclosure, ".//w:rPr/w:sz") == "20"
                and bool(disclosure.xpath(".//w:rPr/w:b", namespaces=ENGINE.NS))
            ),
        })

        formatting_template = root / "formatting-template.docx"
        formatting_output = root / "formatting-output.docx"
        make_formatting_template(formatting_template)
        formatting_inventory = ENGINE.inventory_docx(formatting_template)
        prompt_node = next(node for node in formatting_inventory["nodes"] if node["text"] == "<正文待填>")
        heading_node = next(node for node in formatting_inventory["nodes"] if node["text"] == "实验小结")
        later_node = next(node for node in formatting_inventory["nodes"] if node["text"] == "2.2.2 XXX功能")
        formatting_profile = ENGINE.create_template_profile(
            formatting_inventory,
            [
                {"id": "body", "node_id": prompt_node["node_id"], "operation": "replace_placeholder"},
                {"id": "reflection_title", "node_id": heading_node["node_id"], "operation": "replace_placeholder"},
                {"id": "reflection", "node_id": heading_node["node_id"], "operation": "insert_after"},
                {"id": "later_title", "node_id": later_node["node_id"], "operation": "replace_placeholder"},
            ],
            "格式回归",
        )
        formatting_content = {
            "version": "2.0",
            "items": [
                {"field_id": "body", "content": "已替换正文"},
                {"field_id": "reflection_title", "content": "实验小结（已修正）"},
                {"field_id": "reflection", "content": "第一条正文\n第二条正文"},
                {"field_id": "later_title", "content": "2.2.2 借阅功能"},
            ],
        }
        formatting_result = ENGINE.apply_v2(
            formatting_profile, formatting_template, formatting_content, formatting_output
        )
        with zipfile.ZipFile(formatting_output) as archive:
            formatting_root = ENGINE.etree.fromstring(archive.read("word/document.xml"))
        body_run = formatting_root.xpath('.//w:r[w:t="已替换正文"]', namespaces=ENGINE.NS)[0]
        body_color = ENGINE.child_value(body_run, "./w:rPr/w:color")
        inserted_paragraph = formatting_root.xpath(
            './/w:p[.//w:t="第一条正文"]', namespaces=ENGINE.NS
        )[0]
        inserted_size = ENGINE.child_value(inserted_paragraph, ".//w:rPr/w:sz")
        golden = json.loads(GOLDEN_FORMAT.read_text(encoding="utf-8"))["expected"]
        results.append({
            "name": "docx-semantic-golden",
            "ok": (
                formatting_result["changed_parts"] == golden["changed_parts"]
                and ENGINE.element_text(body_run) == golden["replacement_text"]
                and body_color == golden["replacement_color"]
                and ENGINE.element_text(inserted_paragraph) == golden["inserted_text"]
                and inserted_size == golden["inserted_size_half_points"]
                and ENGINE.active_numbering(inserted_paragraph) is golden["inserted_numbering"]
                and any(golden["remaining_cue_reason"] in cue["reasons"] for cue in formatting_result["remaining_template_cues"])
            ),
        })
        results.append({"name": "placeholder-color-normalized", "ok": body_color == "000000"})
        results.append({
            "name": "insert-after-uses-body-style",
            "ok": not ENGINE.active_numbering(inserted_paragraph) and inserted_size == "24",
        })
        results.append({
            "name": "remaining-template-cues-audited",
            "ok": any(
                "xxx_placeholder" in cue["reasons"]
                for cue in formatting_result["remaining_template_cues"]
            ),
        })
        formatting_text = "\n".join(
            node["text"] for node in ENGINE.inventory_docx(formatting_output)["nodes"]
        )
        results.append({
            "name": "mutation-targets-pre-resolved",
            "ok": "实验小结（已修正）" in formatting_text and "2.2.2 借阅功能" in formatting_text,
        })
        quality_blocked = False
        try:
            ENGINE.apply_v2(
                formatting_profile,
                formatting_template,
                {
                    "version": "2.0",
                    "items": [
                        {
                            "field_id": "body",
                            "content": "太短",
                            "quality": {"unit": "whole", "min_chars": 20, "max_chars": 40},
                        },
                        {"field_id": "reflection_title", "content": "实验小结（已修正）"},
                        {"field_id": "reflection", "content": "内容完整"},
                        {"field_id": "later_title", "content": "2.2.2 借阅功能"},
                    ],
                },
                root / "quality-should-not-write.docx",
            )
        except ENGINE.V2Error as exc:
            quality_blocked = "expected 20-40" in str(exc)
        results.append({"name": "content-quality-gate", "ok": quality_blocked})

        section_template = root / "section-template.docx"
        section_output = root / "section-expanded.docx"
        make_section_template(section_template)
        section_source_hash = ENGINE.file_sha256(section_template)
        section_inventory = ENGINE.inventory_docx(section_template)
        h1 = next(node for node in section_inventory["nodes"] if node["text"] == "1 系统设计")
        h2 = next(node for node in section_inventory["nodes"] if node["text"] == "1.1 功能设计")
        h3 = next(node for node in section_inventory["nodes"] if node["text"] == "1.1.1 登录功能")
        h3_body = next(node for node in section_inventory["nodes"] if node["text"] == "登录功能正文。")
        results.append({
            "name": "inventory-heading-tree",
            "ok": [item["level"] for item in section_inventory["heading_tree"]] == [1, 2, 3],
        })
        section_proposal = {
            "requirements_summary": "任务要求增加测试设计章节，并细分登录测试。",
            "user_request_summary": "用户确认需要扩展模板标题。",
            "material_sources": ["实验任务书.docx", "用户补充说明"],
            "sections": [
                {
                    "id": "test_design", "title": "测试设计", "level": 2,
                    "numbering_mode": "literal", "number_text": "1.2",
                    "style_source_node_id": h2["node_id"],
                    "after_node_id": h3_body["node_id"], "parent_node_id": h1["node_id"],
                },
                {
                    "id": "login_test", "title": "登录测试", "level": 3,
                    "numbering_mode": "literal", "number_text": "1.2.1",
                    "style_source_node_id": h3["node_id"],
                    "after_section_id": "test_design", "parent_section_id": "test_design",
                },
            ],
        }
        section_plan = ENGINE.create_section_plan(section_inventory, section_proposal)
        section_before_entries = xml_entry_hashes(section_template)
        section_result = ENGINE.apply_section_plan(
            section_plan, section_template, section_output, "用户确认新增 1.2 和 1.2.1 标题"
        )
        section_after_entries = xml_entry_hashes(section_output)
        section_output_inventory = ENGINE.inventory_docx(section_output)
        section_texts = [node["text"] for node in section_output_inventory["nodes"]]
        results.append({
            "name": "controlled-section-expansion",
            "ok": (
                "1.2 测试设计" in section_texts
                and "1.2.1 登录测试" in section_texts
                and ENGINE.file_sha256(section_template) == section_source_hash
                and section_result["changed_parts"] == ["word/document.xml"]
                and all(
                    section_before_entries[name] == section_after_entries[name]
                    for name in section_before_entries
                    if name != "word/document.xml"
                )
            ),
        })
        invalid_section_blocked = False
        try:
            ENGINE.create_section_plan(
                section_inventory,
                {
                    **section_proposal,
                    "sections": [{**section_proposal["sections"][0], "level": 3}],
                },
            )
        except ENGINE.V2Error as exc:
            invalid_section_blocked = "style source must be an existing level 3" in str(exc)
        results.append({"name": "section-style-level-mismatch-blocked", "ok": invalid_section_blocked})

        session_workspace = root / "session"
        session = ENGINE.create_session(session_workspace, "软件工程")
        for event in ("confirm_requirements", "resolve_placements", "mark_content_ready", "record_draft"):
            session = ENGINE.advance_session(session_workspace, event, None, None)
        review_id = session["session"]["review_id"]
        session = ENGINE.advance_session(session_workspace, "approve_draft", review_id, "用户确认草稿位置正确")
        session = ENGINE.advance_session(session_workspace, "finalize", review_id, "用户确认生成终稿")
        session = ENGINE.advance_session(session_workspace, "finish_without_update", review_id, "本次不更新 Skill")
        results.append({"name": "session-state-machine", "ok": session["session"]["state"] == "iteration_decided"})
        section_session_workspace = root / "section-session"
        ENGINE.create_session(section_session_workspace, "软件工程")
        ENGINE.advance_session(section_session_workspace, "confirm_requirements", None, "用户确认需求")
        section_state = ENGINE.advance_session(
            section_session_workspace, "resolve_sections", None, "用户确认标题扩展方案"
        )
        section_state = ENGINE.advance_session(
            section_session_workspace, "resolve_placements", None, "用户确认扩展模板定位"
        )
        results.append({
            "name": "optional-section-state-transition",
            "ok": section_state["session"]["state"] == "placements_resolved",
        })

        writing = ENGINE.create_writing_profile("软件工程", "technical", {"tone": "规范"})
        results.append({"name": "writing-profile", "ok": writing["dimensions"]["detail"] == "详细" and writing["dimensions"]["tone"] == "规范"})

        generated = root / "generated.md"
        reference = root / "reference.md"
        repeated = "这是一段长度超过四十个汉字并且不应该从参考实验报告原样复制到最终报告中的测试内容用于验证门禁"
        generated.write_text(repeated + "。", encoding="utf-8")
        reference.write_text(repeated + "。另一段。", encoding="utf-8")
        similarity = ENGINE.similarity_check(generated, [reference], [])
        results.append({"name": "similarity-block", "ok": similarity["gate"] == "blocked"})

        prior_a = root / "prior-a.md"
        prior_b = root / "prior-b.txt"
        prior_a.write_text(
            "# 计算机网络实验\n\n实验目的\n\n"
            "我先配置了本机地址，然后使用抓包工具检查请求。运行后发现第一次解析没有返回结果，"
            "排查后确认是缓存未清理。重新执行命令后得到了预期记录。\n\n"
            "这一步让我理解了缓存对实验现象的影响，也说明观察结果时不能只看一次输出。\n",
            encoding="utf-8",
        )
        prior_b.write_text(
            "软件工程实验中，我先完成接口，再根据测试结果补充异常处理。"
            "最开始遗漏了空输入，测试失败后增加了边界判断。"
            "修改完成后重新运行用例，正常路径和异常路径都得到了对应结果。\n",
            encoding="utf-8",
        )
        personal_profile = ENGINE.analyze_writing_samples([prior_a, prior_b])
        serialized_profile = json.dumps(personal_profile, ensure_ascii=False)
        results.append({
            "name": "prior-report-analysis-extracts-profile-without-body-or-path",
            "ok": (
                personal_profile["kind"] == "personal_writing_profile"
                and personal_profile["sample_summary"]["sample_count"] == 2
                and personal_profile["privacy"]["stores_report_body"] is False
                and str(prior_a) not in serialized_profile
                and "第一次解析没有返回结果" not in serialized_profile
            ),
        })

        ai_heavy = root / "ai-heavy.md"
        ai_heavy.write_text(
            "当然！首先，本实验标志着学习过程中的关键转折点。此外，这不仅仅是一次普通操作，"
            "而是对综合能力的全面提升。专家认为该方法具有十分重要的意义。"
            "综上所述，本实验为后续学习奠定了基础。希望这对您有帮助！\n",
            encoding="utf-8",
        )
        natural = root / "natural.md"
        natural.write_text(
            "配置完成后，我用两组输入检查程序。第一组输出正常，第二组触发了边界错误。"
            "定位到数组长度判断后，我把条件由小于改成小于等于，再次运行时两组结果都符合预期。\n",
            encoding="utf-8",
        )
        ai_audit = ENGINE.humanization_audit(ai_heavy)
        natural_audit = ENGINE.humanization_audit(natural)
        results.append({
            "name": "humanization-audit-uses-clusters-and-keeps-detectors-out-of-gate",
            "ok": (
                ai_audit["status"] == "retryable_failure"
                and natural_audit["status"] == "pass"
                and ai_audit["detector_policy"] == "ai_detector_scores_are_not_quality_gates"
            ),
        })

        cohort_collision = ENGINE.cohort_similarity_check(generated, [reference], [])
        cohort_cold_start = ENGINE.cohort_similarity_check(generated, [], [])
        unrelated = root / "unrelated.md"
        unrelated.write_text(
            "完成接口实现后，测试主要检查空值和重复提交。两个异常分支都返回了预定状态码，"
            "但日志字段仍然不够完整，后面需要补充请求编号。\n",
            encoding="utf-8",
        )
        cohort_clear = ENGINE.cohort_similarity_check(unrelated, [natural], [])
        collision_serialized = json.dumps(cohort_collision, ensure_ascii=False)
        results.append({
            "name": "cohort-collision-gate-is-private-and-distinguishes-clear-output",
            "ok": (
                cohort_collision["gate"] == "blocked"
                and cohort_clear["gate"] == "pass"
                and cohort_cold_start["gate"] == "pass"
                and cohort_cold_start["coverage"]["status"] == "baseline_unavailable"
                and cohort_cold_start["coverage"]["comparison_count"] == 0
                and str(reference) not in collision_serialized
                and repeated not in collision_serialized
            ),
        })

        structured_a = root / "structured-a.docx"
        structured_b = root / "structured-b.docx"
        structured_invalid = root / "structured-invalid.docx"
        problem_without_figure = root / "problem-without-figure.docx"
        problem_with_figure = root / "problem-with-figure.docx"
        make_structured_report(structured_a, "网络地址配置")
        make_structured_report(structured_b, "数据库事务设计")
        make_structured_report(structured_invalid, "类模型设计", valid=False)
        make_problem_evidence_report(problem_without_figure, include_placeholder=False)
        make_problem_evidence_report(problem_with_figure, include_placeholder=True)
        structured_audit = ENGINE.document_structure_audit(structured_a)
        invalid_audit = ENGINE.document_structure_audit(structured_invalid)
        problem_missing_audit = ENGINE.document_structure_audit(problem_without_figure)
        problem_complete_audit = ENGINE.document_structure_audit(problem_with_figure)
        structural_collision = ENGINE.cohort_similarity_check(structured_a, [structured_b], [])
        results.append({
            "name": "document-structure-audit-enforces-numbered-single-placeholders-and-table-decision",
            "ok": (
                structured_audit["status"] == "pass"
                and structured_audit["caption_cross_reference_ok"] is True
                and structured_audit["evidence_counts"] == {
                    "numbered_figures": 1, "numbered_tables": 1, "unnumbered_placeholders": 0,
                }
                and invalid_audit["status"] == "retryable_failure"
                and invalid_audit["single_asset_per_placeholder_ok"] is False
                and bool(invalid_audit["unresolved_template_cues"])
            ),
        })
        results.append({
            "name": "problem-solution-screenshot-placeholder-is-required-only-when-helpful",
            "ok": (
                problem_missing_audit["problem_evidence_needed"] is True
                and problem_missing_audit["problem_evidence_decision_recorded"] is False
                and problem_missing_audit["status"] == "retryable_failure"
                and problem_complete_audit["problem_evidence_decision_recorded"] is True
                and problem_complete_audit["status"] == "pass"
            ),
        })
        results.append({
            "name": "cohort-gate-blocks-same-structure-even-when-topics-differ",
            "ok": (
                structural_collision["gate"] == "blocked"
                and structural_collision["findings"][0]["structural_collision"] is True
                and structural_collision["findings"][0]["structure_flow_similarity"] >= 0.9
            ),
        })

        migration = ENGINE.migrate_v1({"target_document": "old.docx", "items": [{"id": "summary", "target_anchor": "实验小结", "operation": "insert_after"}]})
        results.append({"name": "v1-migration", "ok": migration["fields"][0]["requires_relocation"] is True})

        # Generation-quality contract: only confirmed specs may scaffold a Skill,
        # and the generated package must contain the quality profile and pass the
        # privacy-aware validator.  The frozen binary does not execute source
        # Python helpers, so this check is source-mode only.
        if not getattr(sys, "frozen", False):
            scaffold_script = ROOT / "skills" / "lab-skill-factory" / "scripts" / "scaffold_subject_skill.py"
            validate_skill_script = ROOT / "skills" / "lab-skill-factory" / "scripts" / "validate_scaffolded_skill.py"
            example_spec = ROOT / "skills" / "lab-skill-factory" / "assets" / "skill-spec.example-software-engineering.md"
            confirmed_spec = root / "confirmed-skill-spec.md"
            confirmed_spec.write_text(
                example_spec.read_text(encoding="utf-8").replace(
                    "用户确认情况：待用户确认后才能生成专属 skill",
                    "用户确认情况：已确认（质量回归样例）",
                ),
                encoding="utf-8",
            )
            generated_root = root / "generated-skills"
            generated = subprocess.run(
                [sys.executable, str(scaffold_script), str(confirmed_spec), str(generated_root)],
                capture_output=True,
                text=True,
                check=False,
            )
            generated_path = Path(generated.stdout.strip().splitlines()[-1]) if generated.returncode == 0 and generated.stdout.strip() else None
            quality_validation = subprocess.run(
                [sys.executable, str(validate_skill_script), str(generated_path)] if generated_path else [sys.executable, str(validate_skill_script), str(generated_root / "missing")],
                capture_output=True,
                text=True,
                check=False,
            )
            results.append({
                "name": "skill-quality-scaffold-and-validate",
                "ok": generated.returncode == 0 and quality_validation.returncode == 0 and '"ok": true' in quality_validation.stdout,
                "detail": None if generated.returncode == 0 and quality_validation.returncode == 0 else {
                    "scaffold_stderr": generated.stderr[-1200:],
                    "validation_stdout": quality_validation.stdout[-2400:],
                    "validation_stderr": quality_validation.stderr[-1200:],
                },
            })
            rejected = subprocess.run(
                [sys.executable, str(scaffold_script), str(example_spec), str(root / "rejected-skills")],
                capture_output=True,
                text=True,
                check=False,
            )
            results.append({
                "name": "unconfirmed-spec-blocked",
                "ok": rejected.returncode != 0 and "未确认规则" in (rejected.stdout + rejected.stderr),
            })
        else:
            results.append({"name": "skill-quality-scaffold-and-validate", "ok": True, "skipped": True})
            results.append({"name": "unconfirmed-spec-blocked", "ok": True, "skipped": True})

        # Exercise the public MCP handlers and their workflow preconditions in source mode.
        server_path = ROOT / "mcp" / "lab-skill-factory" / "server.py"
        if server_path.exists():
            server_spec = importlib.util.spec_from_file_location("lab_factory_test_server", server_path)
            assert server_spec and server_spec.loader
            server = importlib.util.module_from_spec(server_spec)
            server_spec.loader.exec_module(server)
            os.environ["LAB_FACTORY_DEV_ALLOW"] = "1"
            api_workspace = root / "api-workspace"
            server.tool_v2_create_session({"workspace": str(api_workspace), "subject": "软件工程"})
            requirements_path = root / "requirements.json"
            ENGINE.write_json(
                requirements_path,
                {
                    "version": "2.0", "experiment": {}, "tools": [], "formatting": {},
                    "evidence": [], "submission": {}, "writable_regions": [],
                    "protected_regions": [], "missing": [], "conflicts": [],
                },
            )
            server.tool_v2_confirm_requirements({"workspace": str(api_workspace), "requirements_path": str(requirements_path), "user_confirmation_summary": "用户确认需求摘要"})
            api_section_workspace = root / "api-section-workspace"
            server.tool_v2_create_session({"workspace": str(api_section_workspace), "subject": "软件工程"})
            server.tool_v2_confirm_requirements({
                "workspace": str(api_section_workspace), "requirements_path": str(requirements_path),
                "user_confirmation_summary": "用户确认需求摘要",
            })
            api_section_inventory_path = root / "api-section-inventory.json"
            server.tool_v2_inventory({
                "docx_path": str(section_template), "output_path": str(api_section_inventory_path),
            })
            api_section_plan_path = root / "api-section-plan.json"
            server.tool_v2_propose_section_plan({
                "inventory_path": str(api_section_inventory_path), "proposal": section_proposal,
                "output_path": str(api_section_plan_path),
            })
            api_expanded_template = root / "api-section-expanded.docx"
            api_section_result = server.tool_v2_apply_section_plan({
                "workspace": str(api_section_workspace), "section_plan_path": str(api_section_plan_path),
                "docx_path": str(section_template), "output_path": str(api_expanded_template),
                "user_confirmation_summary": "用户确认新增测试设计及登录测试标题",
            })
            results.append({
                "name": "mcp-controlled-section-expansion",
                "ok": (
                    api_section_result["session"]["state"] == "sections_resolved"
                    and api_expanded_template.exists()
                ),
            })
            inventory_path = root / "inventory.json"
            server.tool_v2_inventory({"docx_path": str(template), "output_path": str(inventory_path)})
            profile_path = root / "profile.json"
            fields = [
                {"id": "result", "node_id": result_node["node_id"], "operation": "replace_placeholder"},
                {"id": "summary", "node_id": summary_node["node_id"], "operation": "fill_cell"},
            ]
            server.tool_v2_create_template_profile({"inventory_path": str(inventory_path), "subject": "软件工程", "fields": fields, "output_path": str(profile_path)})
            plan_path = root / "placements.json"
            server.tool_v2_propose_placements({"profile_path": str(profile_path), "docx_path": str(template), "output_path": str(plan_path)})
            updated_profile_path = root / "profile-confirmed.json"
            server.tool_v2_resolve_placements({
                "workspace": str(api_workspace), "profile_path": str(profile_path), "docx_path": str(template),
                "placement_plan_path": str(plan_path), "updated_profile_path": str(updated_profile_path),
                "selections": [], "user_confirmation_summary": "全部为高置信度位置",
            })
            server.tool_v2_mark_content_ready({"workspace": str(api_workspace), "content_summary": "两个字段内容已准备"})
            content_path = root / "content.json"
            ENGINE.write_json(content_path, content)
            api_output = root / "api-draft.docx"
            draft_result = server.tool_v2_apply_draft({
                "workspace": str(api_workspace), "profile_path": str(updated_profile_path),
                "docx_path": str(template), "content_package_path": str(content_path), "output_path": str(api_output),
            })
            api_review_id = draft_result["session"]["review_id"]
            server.tool_v2_review_draft({"workspace": str(api_workspace), "review_id": api_review_id, "decision": "approve", "user_feedback_summary": "WPS 检查通过"})
            server.tool_v2_finalize({"workspace": str(api_workspace), "review_id": api_review_id, "generated_path": str(api_output), "reference_paths": [], "user_confirmation_summary": "用户确认终稿"})
            subject_skill = root / "subject-skill"
            subject_skill.mkdir()
            (subject_skill / "SKILL.md").write_text("---\nname: test-skill\ndescription: test\n---\n", encoding="utf-8")
            update_proposal_path = root / "skill-update.json"
            update_proposal = server.tool_v2_propose_skill_update({
                "workspace": str(api_workspace), "skill_dir": str(subject_skill),
                "updates": [{
                    "path": "references/course-rules.json", "scope": "course", "reusable": True,
                    "reason": "保存稳定课程格式", "content": json.dumps({"font": "宋体", "size": "小四"}, ensure_ascii=False),
                }],
                "output_path": str(update_proposal_path),
            })
            update_id = update_proposal["proposal"]["update_id"]
            final_state = server.tool_v2_apply_skill_update({
                "workspace": str(api_workspace), "proposal_path": str(update_proposal_path),
                "update_id": update_id, "review_id": api_review_id,
                "user_confirmation_summary": "用户审阅 diff 后确认更新",
            })
            results.append({"name": "mcp-v2-workflow", "ok": final_state["session"]["state"] == "iteration_decided"})
            try:
                ENGINE.propose_skill_update(subject_skill, [{
                    "path": "references/course-rules.json", "scope": "course", "reusable": True,
                    "content": '{"学号":"20260001"}',
                }])
                privacy_blocked = False
            except ENGINE.V2Error:
                privacy_blocked = True
            results.append({"name": "skill-update-privacy-gate", "ok": privacy_blocked})

    failed = [item["name"] for item in results if not item["ok"]]
    return {"ok": not failed, "count": len(results), "failed": failed, "results": results}


if __name__ == "__main__":
    result = run()
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result["ok"] else 1)
