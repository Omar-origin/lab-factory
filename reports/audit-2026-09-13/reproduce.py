import sys,json,importlib.util,tempfile
from pathlib import Path
from docx import Document
root=Path(__file__).resolve().parents[2]
out=Path(tempfile.mkdtemp(prefix='lab-factory-audit-'))
def module(name,path):
 s=importlib.util.spec_from_file_location(name,path); m=importlib.util.module_from_spec(s); s.loader.exec_module(m); return m
e=module('engine',root/'skills/lab-skill-factory/scripts/v2_engine.py')
t=module('test_auto',root/'mcp/lab-skill-factory/scripts/run_autopilot_tests.py')
def doc(name,paras):
 d=Document()
 for p in paras:d.add_paragraph(p)
 p=out/name;d.save(p);return p
results={}
a=doc('a.docx',['进程调度先依据到达时刻整理就绪队列。','时间片耗尽之后任务重新排到等待队列末尾。','阻塞期间处理器转交给其他可执行任务。','短作业优先策略可能增加长任务的等待时间。','通过记录周转时间比较不同调度策略的表现。','上下文切换需要保存寄存器与程序计数器。','空闲时间段应从处理器利用率计算中扣除。','对于持续到达的新任务需要考虑饥饿现象。'])
b=doc('b.docx',['图像分割首先对彩色照片进行灰度转换。','利用直方图观察亮度分布是否存在明显双峰。','阈值选择会直接影响前景轮廓的完整程度。','中值滤波有助于抑制离散噪点对边缘的干扰。','形态学闭运算能够填补目标内部的小孔洞。','连通区域标记为独立物体分配唯一编号。','面积过小的区域作为噪声候选交由人工复核。','光照不均匀时全局阈值难以覆盖所有局部区域。'])
c=e.cohort_similarity_check(a,[b],[]);results['distinct_prose_same_shape']={'gate':c['gate'],'findings':c['findings']}
d=Document();d.add_paragraph('数据库字段及主键配置如下，详细内容见表 1-1。');d.add_paragraph('表 1-1 数据库字段配置');table=d.add_table(rows=3,cols=2);table.cell(0,0).text='字段';table.cell(0,1).text='主键';table.cell(1,0).text='record_id';table.cell(1,1).text='是';table.cell(2,0).text='label';table.cell(2,1).text='否';p=out/'real-table.docx';d.save(p);v=e.document_structure_audit(p);results['real_table_rejected']={k:v[k] for k in ['status','evidence_counts','table_required_by_content','table_decision_recorded','failures']}
p=doc('placeholder.docx',['本节的界面证据尚未提供，后续需要据实补充。如图 1-1 所示。','【图 1-1：运行界面；待补：真实运行截图】']);v=e.document_structure_audit(p);results['unfilled_evidence_layout_pass']={k:v[k] for k in ['status','failures','evidence_counts']}
art={'draft':{'object_id':'nonexistent-document','sha256':'not-a-valid-hash'},**t.quality_evidence()};results['unbound_quality_attestation']=t.AUTOPILOT.validated_quality_gates(art)
req,prof=t.prepare_files(out);w=out/'session';r=t.AUTOPILOT.prepare(w,t.key(),'审计合成任务','balanced',req,prof,style_identity='audit');c=t.AUTOPILOT.confirm_checkpoint(w,r['session']['state_version'],t.key(),'preflight',r['interaction_plan']['checkpoint']['summary_sha256']);r=t.AUTOPILOT.advance(w,c['session']['state_version'],t.key(),c['confirmation_token'],{});r=t.AUTOPILOT.advance(w,r['session']['state_version'],t.key(),None,art);results['nonexistent_doc_advance']={'state':r['session']['orchestration_state'],'reason':r['reason_code'] if 'reason_code' in r else r['interaction_plan']['reason_code']}
(out/'results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2));print(json.dumps(results,ensure_ascii=False,indent=2))
