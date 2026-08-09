#!/usr/bin/env python3
"""Apply fill.md content to a copied DOCX or Markdown report by fill-map.json.

The script is intentionally conservative:
- the original target document is never modified;
- output starts as a byte-for-byte copy;
- first-pass operations only insert or replace mapped anchors;
- DOCX edits use python-docx and inherit nearby paragraph/table styles where possible.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
import shutil
from pathlib import Path
from typing import Any


SUPPORTED_EXTENSIONS = {".docx", ".md", ".markdown", ".txt"}
DISCLAIMER_TEXT = (
    "AI 辅助生成声明：本文档由 Lab Factory 辅助生成，仅供学习与实验报告草稿参考。"
    "本工具不以实施学术欺诈为目的，不生成或认可伪造的实验数据、截图、运行结果或完成事实。"
    "使用者应核验全部内容、补充真实证据，并遵守所在学校、课程和教师关于 AI 使用及学术诚信的规定。"
    "最终提交与使用责任由使用者承担。"
)


class FillError(Exception):
    pass


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def resolve_path(path_string: str, base_dir: Path) -> Path:
    path = Path(path_string).expanduser()
    if not path.is_absolute():
        path = base_dir / path
    return path.resolve()


def load_fill_map(path: Path) -> dict[str, Any]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise FillError(f"fill-map.json is invalid JSON: {exc}") from exc
    if not isinstance(data, dict):
        raise FillError("fill-map.json top level must be an object")
    if data.get("copy_mode") != "byte_for_byte_first":
        raise FillError("copy_mode must be byte_for_byte_first")
    if data.get("format_strategy") != "inherit_target_anchor":
        raise FillError("format_strategy must be inherit_target_anchor")
    if not isinstance(data.get("items"), list):
        raise FillError("items must be a list")
    for index, item in enumerate(data["items"]):
        if not isinstance(item, dict):
            raise FillError(f"items[{index}] must be an object")
        if item.get("preserve_original") is not True:
            raise FillError(f"items[{index}] preserve_original must be true for first-pass drafts")
        if item.get("format_strategy") != "inherit_target_anchor":
            raise FillError(f"items[{index}] format_strategy must be inherit_target_anchor")
    return data


def normalize_heading(text: str) -> str:
    return re.sub(r"\s+", " ", text.strip().strip("#").strip()).lower()


def parse_markdown_sections(markdown: str) -> dict[str, str]:
    lines = markdown.splitlines()
    headings: list[tuple[int, int, str]] = []
    for index, line in enumerate(lines):
        match = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
        if match:
            headings.append((index, len(match.group(1)), match.group(2).strip()))

    sections: dict[str, str] = {}
    for pos, (line_index, level, title) in enumerate(headings):
        end_index = len(lines)
        for next_line_index, next_level, _ in headings[pos + 1 :]:
            if next_level <= level:
                end_index = next_line_index
                break
        body = "\n".join(lines[line_index + 1 : end_index]).strip()
        sections[normalize_heading(title)] = body
    return sections


def extract_fill_content(source: str, fill_path: Path, fill_sections: dict[str, str]) -> str:
    source = source.strip()
    if source.startswith("literal:"):
        return source[len("literal:") :].strip()
    if "#fill-id:" in source:
        marker = source.split("#fill-id:", 1)[1].strip()
        text = fill_path.read_text(encoding="utf-8", errors="replace")
        pattern = re.compile(
            rf"<!--\s*fill-id:\s*{re.escape(marker)}\s*-->\s*(.*?)(?=<!--\s*fill-id:|^##\s|\Z)",
            re.DOTALL | re.MULTILINE,
        )
        match = pattern.search(text)
        if not match:
            raise FillError(f"fill-id not found in fill.md: {marker}")
        return match.group(1).strip()
    if "#" in source:
        heading = source.rsplit("#", 1)[1]
        key = normalize_heading(heading)
        if key not in fill_sections:
            available = ", ".join(sorted(fill_sections)[:20])
            raise FillError(f"section not found in fill.md: {heading}; available: {available}")
        return fill_sections[key].strip()
    raise FillError(f"unsupported source reference: {source}")


def prepare_output(target_path: Path, output_path: Path, overwrite: bool) -> None:
    if target_path.resolve() == output_path.resolve():
        raise FillError("output_path must not be the same as target_document")
    if output_path.exists() and not overwrite:
        raise FillError(f"output_path already exists: {output_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(target_path, output_path)


def apply_markdown(output_path: Path, items: list[dict[str, Any]], contents: dict[str, str]) -> list[dict[str, Any]]:
    text = output_path.read_text(encoding="utf-8", errors="replace")
    operations = []
    for item in items:
        item_id = str(item.get("id", ""))
        anchor = str(item.get("target_anchor", ""))
        operation = item.get("operation")
        content = contents[item_id]
        if not anchor:
            raise FillError(f"{item_id}: target_anchor is empty")
        if operation == "insert_after":
            lines = text.splitlines()
            found = None
            for index, line in enumerate(lines):
                if anchor in line:
                    found = index
                    break
            if found is None:
                raise FillError(f"{item_id}: anchor not found: {anchor}")
            insertion = content.splitlines() or [content]
            lines[found + 1 : found + 1] = ["", *insertion, ""]
            text = "\n".join(lines) + ("\n" if text.endswith("\n") else "")
            operations.append({"id": item_id, "operation": operation, "anchor": anchor, "status": "applied"})
        elif operation == "replace_placeholder":
            if anchor not in text:
                raise FillError(f"{item_id}: placeholder not found: {anchor}")
            text = text.replace(anchor, content, 1)
            operations.append({"id": item_id, "operation": operation, "anchor": anchor, "status": "applied"})
        elif operation == "fill_table_cell":
            if anchor not in text:
                raise FillError(f"{item_id}: table cell placeholder not found: {anchor}")
            text = text.replace(anchor, content, 1)
            operations.append({"id": item_id, "operation": operation, "anchor": anchor, "status": "applied"})
        else:
            raise FillError(f"{item_id}: unsupported operation for Markdown: {operation}")
    if DISCLAIMER_TEXT not in text:
        text = text.rstrip() + "\n\n---\n\n> " + DISCLAIMER_TEXT + "\n"
    output_path.write_text(text, encoding="utf-8")
    return operations


def copy_run_format(source_run: Any, target_run: Any) -> None:
    target_run.bold = source_run.bold
    target_run.italic = source_run.italic
    target_run.underline = source_run.underline
    target_run.font.name = source_run.font.name
    target_run.font.size = source_run.font.size
    if source_run.font.color and source_run.font.color.rgb:
        target_run.font.color.rgb = source_run.font.color.rgb


def insert_paragraph_after(paragraph: Any, text: str) -> Any:
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph

    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.style = paragraph.style
    new_para.paragraph_format.left_indent = paragraph.paragraph_format.left_indent
    new_para.paragraph_format.first_line_indent = paragraph.paragraph_format.first_line_indent
    new_para.paragraph_format.right_indent = paragraph.paragraph_format.right_indent
    new_para.paragraph_format.space_before = paragraph.paragraph_format.space_before
    new_para.paragraph_format.space_after = paragraph.paragraph_format.space_after
    new_para.paragraph_format.line_spacing = paragraph.paragraph_format.line_spacing
    run = new_para.add_run(text)
    if paragraph.runs:
        copy_run_format(paragraph.runs[0], run)
    return new_para


def all_paragraphs(document: Any) -> list[Any]:
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    return paragraphs


def replace_in_paragraph(paragraph: Any, anchor: str, content: str) -> bool:
    if anchor not in paragraph.text:
        return False
    for run in paragraph.runs:
        if anchor in run.text:
            run.text = run.text.replace(anchor, content, 1)
            return True
    # Anchor may span runs. Use the first run and clear the rest to avoid rebuilding the whole document.
    if paragraph.runs:
        original = paragraph.text
        paragraph.runs[0].text = original.replace(anchor, content, 1)
        for run in paragraph.runs[1:]:
            run.text = ""
        return True
    paragraph.add_run(content)
    return True


def parse_table_anchor(anchor: str) -> tuple[int, int, int] | None:
    match = re.match(r"^table:(\d+)[:;, ]+row:(\d+)[:;, ]+col:(\d+)$", anchor.strip(), re.I)
    if match:
        return tuple(int(part) for part in match.groups())  # type: ignore[return-value]
    match = re.match(r"^table:(\d+)[:;, ]+(\d+)[:;, ]+(\d+)$", anchor.strip(), re.I)
    if match:
        return tuple(int(part) for part in match.groups())  # type: ignore[return-value]
    return None


def apply_docx(output_path: Path, items: list[dict[str, Any]], contents: dict[str, str]) -> list[dict[str, Any]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise FillError("python-docx is not available; cannot write DOCX") from exc

    document = Document(str(output_path))
    operations = []
    for item in items:
        item_id = str(item.get("id", ""))
        anchor = str(item.get("target_anchor", ""))
        operation = item.get("operation")
        content = contents[item_id]
        if not anchor:
            raise FillError(f"{item_id}: target_anchor is empty")

        if operation == "insert_after":
            matched = None
            for paragraph in all_paragraphs(document):
                if anchor in paragraph.text:
                    matched = paragraph
                    break
            if matched is None:
                raise FillError(f"{item_id}: anchor not found: {anchor}")
            current = matched
            for line in reversed([line for line in content.splitlines() if line.strip()]):
                current = insert_paragraph_after(matched, line)
            operations.append({"id": item_id, "operation": operation, "anchor": anchor, "status": "applied"})
        elif operation == "replace_placeholder":
            applied = False
            for paragraph in all_paragraphs(document):
                if replace_in_paragraph(paragraph, anchor, content):
                    applied = True
                    break
            if not applied:
                raise FillError(f"{item_id}: placeholder not found: {anchor}")
            operations.append({"id": item_id, "operation": operation, "anchor": anchor, "status": "applied"})
        elif operation == "fill_table_cell":
            parsed = parse_table_anchor(anchor)
            if parsed:
                table_index, row_index, col_index = parsed
                try:
                    cell = document.tables[table_index].rows[row_index].cells[col_index]
                except IndexError as exc:
                    raise FillError(f"{item_id}: table anchor out of range: {anchor}") from exc
                if cell.paragraphs:
                    paragraph = cell.paragraphs[0]
                    paragraph.text = ""
                    run = paragraph.add_run(content)
                    if len(cell.paragraphs) > 1 and cell.paragraphs[1].runs:
                        copy_run_format(cell.paragraphs[1].runs[0], run)
                else:
                    cell.text = content
                operations.append({"id": item_id, "operation": operation, "anchor": anchor, "status": "applied"})
            else:
                applied = False
                for table in document.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            for paragraph in cell.paragraphs:
                                if replace_in_paragraph(paragraph, anchor, content):
                                    applied = True
                                    break
                            if applied:
                                break
                        if applied:
                            break
                    if applied:
                        break
                if not applied:
                    raise FillError(f"{item_id}: table placeholder not found: {anchor}")
                operations.append({"id": item_id, "operation": operation, "anchor": anchor, "status": "applied"})
        else:
            raise FillError(f"{item_id}: unsupported operation for DOCX: {operation}")

    matching = [paragraph for paragraph in document.paragraphs if paragraph.text.strip() == DISCLAIMER_TEXT]
    for duplicate in matching[1:]:
        duplicate._element.getparent().remove(duplicate._element)
    if not matching:
        document.add_paragraph(DISCLAIMER_TEXT)
    document.save(str(output_path))
    return operations


def apply_fill_map(fill_map_path: Path, output_path: Path | None, overwrite: bool) -> dict[str, Any]:
    fill_map = load_fill_map(fill_map_path)
    base_dir = fill_map_path.parent
    fill_path = resolve_path(str(fill_map["source_fill"]), base_dir)
    target_path = resolve_path(str(fill_map["target_document"]), base_dir)
    if not fill_path.exists():
        raise FillError(f"source_fill does not exist: {fill_path}")
    if not target_path.exists():
        raise FillError(f"target_document does not exist: {target_path}")
    suffix = target_path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise FillError(f"unsupported target extension: {suffix}")

    if output_path is None:
        output_path = target_path.with_name(f"{target_path.stem}-draft{target_path.suffix}")
    else:
        output_path = output_path.expanduser().resolve()

    source_hash_before = sha256_file(target_path)
    prepare_output(target_path, output_path, overwrite)
    copied_hash = sha256_file(output_path)
    if copied_hash != source_hash_before:
        raise FillError("byte-for-byte copy check failed before editing")

    fill_text = fill_path.read_text(encoding="utf-8", errors="replace")
    fill_sections = parse_markdown_sections(fill_text)
    contents: dict[str, str] = {}
    for item in fill_map["items"]:
        item_id = str(item.get("id", ""))
        if not item_id:
            raise FillError("each fill-map item must have id")
        contents[item_id] = extract_fill_content(str(item.get("source", "")), fill_path, fill_sections)

    if suffix == ".docx":
        operations = apply_docx(output_path, fill_map["items"], contents)
    else:
        operations = apply_markdown(output_path, fill_map["items"], contents)

    source_hash_after = sha256_file(target_path)
    if source_hash_after != source_hash_before:
        raise FillError("source document changed; aborting because source files must stay read-only")

    return {
        "ok": True,
        "fill_map": str(fill_map_path),
        "source_fill": str(fill_path),
        "target_document": str(target_path),
        "output_document": str(output_path),
        "target_sha256": source_hash_before,
        "output_sha256": sha256_file(output_path),
        "operation_count": len(operations),
        "operations": operations,
        "limitations": [
            "DOCX complex objects such as text boxes, formulas, floating images, fields, comments, and tracked changes may require manual handling.",
            "First-pass output only fills mapped anchors; deleting template hints must wait for user-confirmed finalize.",
        ],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Apply fill-map.json to a copied DOCX or Markdown document.")
    parser.add_argument("fill_map", help="Path to fill-map.json.")
    parser.add_argument("--output", help="Output draft document path. Defaults to <target>-draft.<ext>.")
    parser.add_argument("--overwrite", action="store_true", help="Allow overwriting output path.")
    args = parser.parse_args()

    try:
        result = apply_fill_map(
            Path(args.fill_map).expanduser().resolve(),
            Path(args.output).expanduser().resolve() if args.output else None,
            args.overwrite,
        )
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except FillError as exc:
        print(json.dumps({"ok": False, "error": str(exc)}, ensure_ascii=False, indent=2))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
